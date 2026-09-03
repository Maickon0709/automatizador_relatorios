"""
test_relatorio.py

Testes do módulo relatorio.py: geração do arquivo Excel (3 abas) e
da imagem-resumo (.png). Usa uma pasta temporária (tmp_path) para não
sujar a pasta do projeto com arquivos gerados durante os testes.
"""

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from processador import processar
from relatorio import gerar_relatorio, gerar_imagem_relatorio


@pytest.fixture
def resultado_processado():
    df_vendas = pd.DataFrame({
        "data": pd.to_datetime(["2026-08-01", "2026-08-02"]),
        "vendedor": ["João Silva", "Maria Souza"],
        "numero_produto": ["PRD-1023", "PRD-1050"],
        "nome_produto": ["Cadeira Gamer", "Mouse Sem Fio"],
        "quantidade_vendida": [3, 5],
        "valor_unitario": [450.0, 90.0],
    })
    return processar(df_vendas)


def test_gerar_relatorio_cria_arquivo_xlsx(resultado_processado, tmp_path):
    caminho = gerar_relatorio(resultado_processado, pasta_saida=str(tmp_path))

    assert Path(caminho).exists()
    assert caminho.endswith(".xlsx")


def test_gerar_relatorio_tem_tres_abas(resultado_processado, tmp_path):
    caminho = gerar_relatorio(resultado_processado, pasta_saida=str(tmp_path))
    planilha = load_workbook(caminho)

    assert set(planilha.sheetnames) == {
        "Vendas Detalhadas", "Resumo por Vendedor", "Resumo por Produto"
    }


def test_gerar_relatorio_chave_faltando_gera_erro(tmp_path):
    resultado_incompleto = {"vendas_detalhadas": pd.DataFrame()}

    with pytest.raises(ValueError):
        gerar_relatorio(resultado_incompleto, pasta_saida=str(tmp_path))


def test_gerar_imagem_relatorio_cria_arquivo_png(resultado_processado, tmp_path):
    caminho = gerar_imagem_relatorio(resultado_processado, pasta_saida=str(tmp_path))

    assert Path(caminho).exists()
    assert caminho.endswith(".png")
    assert Path(caminho).stat().st_size > 0  # não é um arquivo vazio/corrompido
